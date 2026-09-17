"""Read workbooks without modifying or recalculating them.

The JSON representation is also the contract for a future live Sheets adapter.
Formula caches are observations, not proof of a recent calculation.
"""
from __future__ import annotations

import datetime as dt
import math
import warnings
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.cell.cell import Cell
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.utils.cell import range_boundaries, get_column_letter

STYLE_FIELDS = ("font", "fill", "border", "alignment", "protection", "number_format")


def scalar(value):
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return {"date": value.isoformat()}
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError("Workbook contains a non-finite numeric value")
    return value


def theme_colors(workbook):
    if not workbook.loaded_theme:
        return []
    root = ET.fromstring(workbook.loaded_theme)
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    scheme = root.find(".//a:clrScheme", ns)
    colors = {}
    if scheme is not None:
        for node in scheme:
            color = next(iter(node))
            colors[node.tag.split("}")[-1]] = color.get("lastClr", color.get("val"))
    return [colors.get(name) for name in
            ("lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3", "accent4",
             "accent5", "accent6", "hlink", "folHlink")]


def xml(obj, remove=(), palette=()):
    if obj is None:
        return None
    root = obj if ET.iselement(obj) else obj.to_tree()
    axis_ids = {}
    for node in root.iter():
        if node.tag.split('}')[-1] == 'axId' and node.get('val') not in axis_ids:
            axis_ids[node.get('val')] = str(len(axis_ids) + 1)
    for node in root.iter():
        if node.tag.split('}')[-1] in ('axId', 'crossAx') and node.get('val') in axis_ids:
            node.set('val', axis_ids[node.get('val')])
        for key in remove:
            node.attrib.pop(key, None)
        if node.tag.split("}")[-1] in ("color", "fgColor", "bgColor"):
            color = node.get("rgb")
            if "theme" in node.attrib:
                index = int(node.get("theme"))
                color = palette[index] if index < len(palette) else None
            elif "indexed" in node.attrib:
                index = int(node.get("indexed"))
                color = COLOR_INDEX[index] if index < len(COLOR_INDEX) else ("000000" if index == 64 else "FFFFFF")
            elif node.get("auto") == "1":
                color = "000000"
            if color:
                tint = node.get("tint")
                node.attrib.clear()
                node.set("rgb", color[-6:].upper())
                if tint and float(tint) != 0:
                    node.set("tint", tint)
    # Chart cached series values change after calculation without an object edit.
    for parent in root.iter():
        for child in list(parent):
            if child.tag.split("}")[-1] in ("numCache", "strCache", "multiLvlStrCache"):
                parent.remove(child)
    return ET.tostring(root, encoding="unicode")


def style(cell, palette=()):
    return {key: cell.number_format if key == "number_format" else xml(getattr(cell, key), palette=palette)
            for key in STYLE_FIELDS}


def formula(cell):
    if cell.data_type != "f":
        return None
    if isinstance(cell.value, str):
        return cell.value
    # Array formulas must be compared by their contents, never Python object identity.
    if hasattr(cell.value, "text"):
        return cell.value.text
    return {"kind": type(cell.value).__name__, "attributes": dict(cell.value)}


def raw(cell):
    value = cell.get("value")
    kind = ("boolean" if isinstance(value, bool) else "number" if isinstance(value, (int, float))
            else "blank" if value is None else "text" if isinstance(value, str) else "date")
    return {"formula": cell["formula"]} if cell.get("formula") is not None else {
        "value": value, "kind": kind, "error": cell.get("error")}


def coordinates(ref):
    lo_col, lo_row, hi_col, hi_row = range_boundaries(ref)
    if None in (lo_col, lo_row, hi_col, hi_row):
        raise ValueError(f"Use bounded A1 ranges, not entire rows/columns: {ref}")
    for row in range(lo_row, hi_row + 1):
        for col in range(lo_col, hi_col + 1):
            yield f"{get_column_letter(col)}{row}"


def expand(ranges):
    return {(sheet, address) for sheet, refs in ranges.items()
            for ref in refs for address in coordinates(ref)}


def get_cell(book, sheet, address):
    return book.get("sheets", {}).get(sheet, {}).get("cells", {}).get(address, {
        "value": None, "formula": None, "error": None})


def all_cells(book):
    return {(sheet, address) for sheet, data in book["sheets"].items()
            for address in data["cells"]}


def load_workbook(path):
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        raise ValueError("Submission must be .xlsx or a grader JSON snapshot")
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        formulas = openpyxl.load_workbook(path, data_only=False)
        values = openpyxl.load_workbook(path, data_only=True)
    result = {
        "schema_version": 1,
        "calculation": "unverified_cache",
        "source_warnings": sorted({str(w.message) for w in caught}),
        "sheet_order": formulas.sheetnames,
        "defined_names": sorted(xml(n) for n in formulas.defined_names.values()),
        "sheets": {},
    }
    palette = theme_colors(formulas)
    for sheet in formulas:
        defaults = style(Cell(sheet, row=1, column=1), palette)
        style_cache = {}
        cells = {}
        # Read allocated cells only; avoid materializing huge rectangular blank ranges.
        for cell in sheet._cells.values():
            key = str(cell._style)
            if key not in style_cache:
                style_cache[key] = style(cell, palette)
            cell_style = style_cache[key]
            if cell.value is None and cell_style == defaults:
                continue
            cached = values[sheet.title][cell.coordinate]
            f = formula(cell)
            cells[cell.coordinate] = {
                "formula": f,
                "value": scalar(cached.value if f is not None else cell.value),
                "error": cached.value if cached.data_type == "e" else None,
                "style": cell_style,
            }
        conditional_formats = []
        for cf, rules in sheet.conditional_formatting._cf_rules.items():
            for rule in rules:
                conditional_formats.append({"range": str(cf.sqref),
                    "rule": xml(rule, remove=("dxfId",), palette=palette), "style": xml(rule.dxf, palette=palette)})
        result["sheets"][sheet.title] = {
            "cells": cells,
            "default_style": defaults,
            "structure": {
                "state": sheet.sheet_state,
                "merged_ranges": sorted(str(r) for r in sheet.merged_cells.ranges),
                "row_dimensions": {str(k): {field: getattr(v, field) for field in
                    ("height", "hidden", "outlineLevel", "collapsed")}
                    for k, v in sheet.row_dimensions.items()},
                "column_dimensions": {str(k): {field: getattr(v, field) for field in
                    ("min", "max", "width", "hidden", "outlineLevel", "collapsed")}
                    for k, v in sheet.column_dimensions.items()},
            },
            "objects": {
                "validations": [xml(v) for v in sheet.data_validations.dataValidation],
                "conditional_formats": conditional_formats,
                "charts": [xml(c._write(), palette=palette) for c in sheet._charts],
                "pivots": [xml(p) for p in sheet._pivots],
                "tables": [xml(t) for t in sheet.tables.values()],
            },
        }
    formulas.close()
    values.close()
    return result
